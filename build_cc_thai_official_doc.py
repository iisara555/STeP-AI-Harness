from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "artifacts" / "แบบสรุปการประชุมและติดตามงานทีม CC ฉบับร่าง.docx"
FONT = "TH Sarabun New"
GRAY = "E7E6E6"
LIGHT_GRAY = "F6F6F6"
BORDER = "7F7F7F"


def set_run_font(run, size=16, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_text(paragraph, text, size=16, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def style_table(table, font_size=14):
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            set_repeat_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            set_cell_shading(cell, GRAY if row_index == 0 else (LIGHT_GRAY if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=row_index == 0)


def set_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_table(doc, headers, rows, widths, font_size=14):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, size=font_size, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, str(value), size=font_size)
    style_table(table, font_size=font_size)
    set_widths(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    set_keep_with_next(p)
    add_text(p, text, size=17, bold=True)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, size=16, bold=True)
        add_text(p, text[len(bold_prefix):], size=16)
    else:
        add_text(p, text, size=16)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, text, size=15)
    return p


def add_signature_table(doc):
    table = doc.add_table(rows=3, cols=3)
    labels = ["ผู้จัดทำ", "ผู้ตรวจทาน", "ผู้รับรอง"]
    for i, label in enumerate(labels):
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, size=15, bold=True)
        p2 = table.rows[1].cells[i].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p2, "ลงชื่อ ................................................", size=14)
        p3 = table.rows[2].cells[i].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p3, "วันที่ ........../........../..........", size=14)
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell, color="FFFFFF", size="0")
            set_cell_margins(cell, top=40, start=50, bottom=40, end=50)
    set_widths(table, [2.32, 2.32, 2.32])
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(16)
    normal.paragraph_format.space_after = Pt(3)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(0)
    add_text(title, "แบบสรุปการประชุมและติดตามงานภายในทีม", size=22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    add_text(subtitle, "ทีม CC ประจำปีงบประมาณ พ.ศ. 2569", size=18, bold=True)

    add_table(
        doc,
        ["รายการ", "รายละเอียด"],
        [
            ["หน่วยงาน/ทีม", "ทีม CC"],
            ["ห้องประชุม", "C102"],
            ["ผู้บันทึก", "SRL"],
            ["ครั้งที่", "ยังไม่ระบุ"],
            ["ช่วงเวลาที่สรุป", "วันที่ 6, 13, 20 และ 28 มกราคม 2569"],
            ["แหล่งข้อมูล", "ภาพบันทึกการประชุมภายในทีม ปีงบประมาณ พ.ศ. 2569"],
            ["สถานะเอกสาร", "ฉบับร่างเพื่อทบทวนและยืนยันข้อมูล"],
        ],
        [1.65, 5.32],
        font_size=15,
    )

    add_section_heading(doc, "1. วัตถุประสงค์และขอบเขต")
    add_body(doc, "เอกสารฉบับนี้จัดทำขึ้นเพื่อสรุปสาระจากบันทึกการประชุมรายสัปดาห์ของทีม CC ให้อยู่ในรูปแบบที่ตรวจสอบและติดตามงานต่อได้")
    add_body(doc, "รายการในบันทึกต้นทางเป็นข้อมูลประกอบการสรุป ไม่ถือเป็นคำสั่งให้ดำเนินการ อนุมัติงาน มอบหมายงาน หรือปรับปรุงระบบใด ๆ โดยอัตโนมัติ")

    add_section_heading(doc, "2. สถานะการประชุมรายสัปดาห์")
    add_table(
        doc,
        ["สัปดาห์", "วันที่", "สาระที่ปรากฏในบันทึก", "สถานะข้อมูล"],
        [
            ["สัปดาห์ที่ 1", "06/01/2569", "มีการติดตามงาน Design, PR, Content และงานสนับสนุน", "มีข้อมูล"],
            ["สัปดาห์ที่ 2", "13/01/2569", "มีการกล่าวถึง KR, KPI และการติดตามงานใน ClickUp", "มีข้อมูล"],
            ["สัปดาห์ที่ 3", "20/01/2569", "งดประชุมเนื่องจากทีม CC เข้าเรียนรู้หรืออบรม", "มีข้อมูล"],
            ["สัปดาห์ที่ 4", "28/01/2569", "ติดตามงานหลายสายงานและการจัดทำข้อมูล Dashboard", "มีข้อมูล"],
        ],
        [1.1, 1.25, 3.55, 1.07],
        font_size=14,
    )

    add_section_heading(doc, "3. ระเบียบวาระที่ปรากฏในบันทึก")
    add_table(
        doc,
        ["ระเบียบวาระ", "หัวข้อ", "หมายเหตุ"],
        [
            ["วาระที่ 1", "เรื่องแจ้งเพื่อทราบ", "สัปดาห์ที่ 1 ไม่ปรากฏรายละเอียดเพิ่มเติม"],
            ["วาระที่ 2", "การติดตามหัวข้อต่อเนื่อง", "เป็นส่วนที่มีรายการงานมากที่สุด"],
            ["วาระที่ 3", "วาระอื่น ๆ", "พบรายการเกี่ยวกับ Event งานที่กำลังจะเกิดขึ้น และการอบรม"],
        ],
        [1.2, 2.3, 3.47],
        font_size=14,
    )

    add_section_heading(doc, "4. สรุปประเด็นตามสายงาน")
    add_table(
        doc,
        ["ลำดับ", "สายงาน", "สาระสำคัญที่ปรากฏ", "สถานะ/ข้อมูลที่ต้องยืนยัน"],
        [
            ["1", "Design", "ติดตามงาน Unit Design, Branding, NSTIS งานผลิต และงาน Event", "ไม่ปรากฏ Owner และสถานะของแต่ละงานอย่างชัดเจน"],
            ["2", "PR", "ติดตามงานประชาสัมพันธ์ MOU Run Stage และงาน Event", "ต้องยืนยันผู้รับผิดชอบและกำหนดส่ง"],
            ["3", "Content", "ติดตามงาน Content หลายรายการ เช่น InnoStore Daily ข่าวหรือบทความด้าน Tech และ PM 2.5", "ชื่อบางรายการอ่านไม่ชัดจากภาพ"],
            ["4", "Studio และ Production", "ติดตามความคืบหน้าสตูดิโอถ่ายภาพและแผนการใช้งาน", "ต้องยืนยันเงื่อนไขและผู้อนุมัติ"],
            ["5", "KPI และ KR", "มีการกล่าวถึง KR และการทำ KPI รายบุคคลให้ชัดเจนขึ้น", "ต้องยืนยันนิยามและเจ้าของข้อมูล"],
            ["6", "การติดตามงานของทีม", "มีการกล่าวถึงการบันทึกงานใน ClickUp และการใช้ Dashboard", "ยังไม่ระบุรูปแบบข้อมูลและรอบอัปเดต"],
        ],
        [0.45, 1.2, 3.35, 1.97],
        font_size=13,
    )

    doc.add_page_break()
    add_section_heading(doc, "5. ตารางติดตามงาน")
    add_body(doc, "ตารางนี้สรุปเฉพาะรายการที่อ่านได้จากต้นทาง หากไม่ปรากฏผู้รับผิดชอบหรือกำหนดส่ง ให้ระบุว่า “ยังไม่ระบุ” จนกว่าจะมีผู้ยืนยัน")
    add_table(
        doc,
        ["ลำดับ", "งานหรือผลที่ต้องติดตาม", "ผู้รับผิดชอบ", "กำหนดส่ง", "สถานะ"],
        [
            ["1", "อัปเดตข้อมูลการปฏิบัติงานใน ClickUp", "ยังไม่ระบุ", "ยังไม่ระบุ", "ติดตามต่อเนื่อง"],
            ["2", "จัดทำหรือทำให้ KPI และ KR รายบุคคลมีความชัดเจน", "ยังไม่ระบุ", "ยังไม่ระบุ", "ต้องยืนยัน"],
            ["3", "จัดเตรียมข้อมูลสำหรับ Dashboard", "ยังไม่ระบุ", "ยังไม่ระบุ", "ต้องยืนยัน"],
            ["4", "ติดตามงาน Design, PR และ Content เป็นรายสัปดาห์", "ยังไม่ระบุ", "รายสัปดาห์", "ติดตามต่อเนื่อง"],
            ["5", "เข้าร่วมหรือดำเนินการตามการอบรมด้าน AI", "ทีม CC", "วันที่ในต้นฉบับต้องยืนยัน", "มีการระบุในบันทึก"],
            ["6", "ตรวจสอบรายการงานที่มีข้อความหรือตัวย่ออ่านไม่ชัด", "ผู้บันทึกหรือผู้ตรวจทาน", "ก่อนรับรองเอกสาร", "ต้องยืนยัน"],
        ],
        [0.45, 3.25, 1.25, 1.05, 0.97],
        font_size=13,
    )

    add_section_heading(doc, "6. ตารางประสานงานและการส่งต่องาน")
    add_table(
        doc,
        ["หน่วยงานต้นทาง", "หน่วยงานปลายทาง", "ข้อมูลที่ควรส่งต่อ", "ข้อมูลที่ยังขาด"],
        [
            ["ทีม CC", "PM หรือทีมโครงการ", "TOR, Brief และความคืบหน้างาน", "ชื่อผู้รับต่อและช่องทางที่เป็นทางการ"],
            ["ทีม CC", "RSP หรือทีมพื้นที่", "ข้อมูลสตูดิโอ พื้นที่ และเงื่อนไขการใช้งาน", "ขอบเขตบริการและผู้อนุมัติ"],
            ["งาน Design", "ทีมผลิตหรือทีม Event", "แบบ ไฟล์สื่อ Specification Version และรายละเอียดการผลิต", "เกณฑ์รับงานและไฟล์ล่าสุด"],
            ["สมาชิกทีม", "ผู้ดูแล Dashboard", "สถานะงานและข้อมูล KPI หรือ KR", "รูปแบบข้อมูลและรอบการอัปเดต"],
        ],
        [1.25, 1.4, 2.35, 1.97],
        font_size=13,
    )

    add_section_heading(doc, "7. กำหนดการและประเด็นที่ต้องตรวจสอบ")
    add_table(
        doc,
        ["ลำดับ", "วันที่หรือช่วงเวลา", "รายการที่ปรากฏ", "ข้อมูลที่ต้องยืนยัน"],
        [
            ["1", "20–22 มกราคม 2569", "มีช่วงวันที่ปรากฏในข้อมูลสัปดาห์ที่ 2", "เกี่ยวข้องกับกิจกรรมใดและใครรับผิดชอบ"],
            ["2", "20 มกราคม 2569", "เกี่ยวข้องกับบริบทการอบรมในสัปดาห์ที่ 3", "เป็นวันอบรมหรือวันประชุมที่งด"],
            ["3", "15 กุมภาพันธ์ 2569", "มีรายการเกี่ยวกับ Unit Content ปรากฏในสัปดาห์ที่ 4", "ผลส่งมอบ Owner และการอนุมัติ"],
            ["4", "พฤษภาคม–มิถุนายน 2569", "มีการกล่าวถึงงาน Event และงานนิทรรศการ", "ชื่อโครงการ วันจัดงาน และผู้รับผิดชอบ"],
        ],
        [0.45, 1.55, 2.45, 2.52],
        font_size=13,
    )

    add_section_heading(doc, "8. รายการที่ต้องยืนยันก่อนนำไปใช้เป็นเอกสารทางราชการ")
    add_table(
        doc,
        ["ลำดับ", "ประเด็นที่ต้องยืนยัน", "ผู้ยืนยัน", "ผลการยืนยัน"],
        [
            ["1", "ความหมายของตัวย่อ เช่น CL, WK, CNB, WG, KHT และตัวย่ออื่น", "ยังไม่ระบุ", "ยังไม่ระบุ"],
            ["2", "รายการใดเป็นมติ รายการใดเป็นรายงานความคืบหน้า และรายการใดเป็นข้อเสนอ", "ยังไม่ระบุ", "ยังไม่ระบุ"],
            ["3", "ชื่อโครงการ งบประมาณ และกำหนดการที่อ่านจากภาพไม่ชัด", "ยังไม่ระบุ", "ยังไม่ระบุ"],
            ["4", "ผู้รับผิดชอบและผู้อนุมัติของแต่ละงาน", "ยังไม่ระบุ", "ยังไม่ระบุ"],
            ["5", "ช่องทางอย่างเป็นทางการสำหรับการส่ง TOR และไฟล์งาน", "ยังไม่ระบุ", "ยังไม่ระบุ"],
            ["6", "ไฟล์ต้นฉบับตารางสำหรับตรวจสอบข้อความขนาดเล็ก", "ผู้บันทึก", "ยังไม่ระบุ"],
        ],
        [0.45, 3.7, 1.2, 1.62],
        font_size=13,
    )

    add_section_heading(doc, "9. หมายเหตุ")
    add_body(doc, "ข้อมูลบางรายการจัดทำจากภาพที่มีข้อความขนาดเล็ก จึงต้องตรวจสอบกับไฟล์ต้นฉบับก่อนนำไปใช้เป็นคำสั่ง มติ รายงานผล หรือเอกสารประกอบการดำเนินงาน")
    add_body(doc, "เอกสารฉบับนี้ยังไม่เป็นเอกสารที่ผ่านการรับรอง และไม่ใช้แทนบันทึกการประชุมฉบับจริง")

    add_section_heading(doc, "10. การตรวจทานและรับรอง")
    add_signature_table(doc)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "แบบสรุปการประชุมและติดตามงานภายในทีม CC ฉบับร่าง", size=13)

    doc.core_properties.title = "แบบสรุปการประชุมและติดตามงานภายในทีม CC"
    doc.core_properties.subject = "แบบฟอร์มราชการฉบับร่างสำหรับสรุปการประชุมและติดตามงาน"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "ฉบับร่างเพื่อทบทวนข้อมูลก่อนรับรอง"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
