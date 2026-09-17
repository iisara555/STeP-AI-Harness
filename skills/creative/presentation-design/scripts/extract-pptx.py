#!/usr/bin/env python3
"""
STeP Presentation Design — PowerPoint Content & Asset Extractor
สกัดข้อความ โครงสร้างสไลด์ รูปภาพ และ Speaker Notes จากไฟล์ .pptx

การใช้งาน:
    python extract-pptx.py <input.pptx> [output_dir]

ความต้องการ:
    pip install python-pptx
"""

import json
import os
import sys

def extract_pptx(file_path, output_dir="."):
    try:
        from pptx import Presentation
    except ImportError:
        print("❌ ต้องติดตั้ง python-pptx ก่อน: pip install python-pptx")
        sys.exit(1)

    prs = Presentation(file_path)
    slides_data = []

    assets_dir = os.path.join(output_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)

    for slide_num, slide in enumerate(prs.slides):
        slide_data = {
            "number": slide_num + 1,
            "title": "",
            "content": [],
            "images": [],
            "notes": "",
        }

        for shape in slide.shapes:
            if shape.has_text_frame:
                if shape == slide.shapes.title:
                    slide_data["title"] = shape.text.strip()
                else:
                    text = shape.text.strip()
                    if text:
                        slide_data["content"].append({"type": "text", "content": text})

            # Extract Pictures (Shape Type 13 = Picture)
            if shape.shape_type == 13:
                try:
                    image = shape.image
                    image_bytes = image.blob
                    image_ext = image.ext
                    image_name = f"slide{slide_num + 1}_img{len(slide_data['images']) + 1}.{image_ext}"
                    image_path = os.path.join(assets_dir, image_name)

                    with open(image_path, "wb") as f:
                        f.write(image_bytes)

                    slide_data["images"].append({
                        "path": f"assets/{image_name}",
                        "width": shape.width,
                        "height": shape.height,
                    })
                except (IOError, AttributeError, ValueError) as e:
                    print(f"⚠️ ไม่สามารถบันทึกภาพในสไลด์ {slide_num + 1}: {e}")

        # Speaker notes
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            slide_data["notes"] = notes_frame.text.strip()

        slides_data.append(slide_data)

    output_path = os.path.join(output_dir, "extracted-slides.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(slides_data, f, ensure_ascii=False, indent=2)

    print(f"✅ สกัดเนื้อหาสำเร็จ {len(slides_data)} สไลด์ -> บันทึกที่ {output_path}")
    for s in slides_data:
        imgs = len(s["images"])
        print(f"  สไลด์ {s['number']}: {s['title'] or '(ไม่มีชื่อเรื่อง)'} — รูป {imgs} ภาพ")

    return slides_data

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("วิธีใช้: python extract-pptx.py <ไฟล์.pptx> [โฟลเดอร์ปลายทาง]")
        sys.exit(1)

    input_file = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else "."
    extract_pptx(input_file, out_dir)
