from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "artifacts" / "cc-team-weekly-work-review-draft.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_font(run, name="Aptos", size=10, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, bold=False, color="000000", size=10):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def style_table(table, header_fill="1F4E78"):
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F6F9")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_font(run, size=8.5, bold=row_index == 0, color="FFFFFF" if row_index == 0 else "000000")


def set_table_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(paragraph, header, bold=True, color="FFFFFF", size=8.5)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            paragraph = cells[i].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_text(paragraph, str(value), size=8.5)
    style_table(table)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    set_keep_with_next(paragraph)
    add_text(paragraph, text, bold=True, size=13 if level == 1 else 10.5)
    return paragraph


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.08
    add_text(paragraph, text, size=9.5)
    return paragraph


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    add_text(title, "CC Team Weekly Work Review", bold=True, size=20)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    add_text(subtitle, "Draft extraction from the weekly internal meeting record", color="5B6573", size=10.5)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(7)
    add_text(intro, "Purpose. ", bold=True, size=9.5)
    add_text(
        intro,
        "Convert the supplied CC team meeting image into a reviewable record of workstreams, actions, handoffs, and items that require human confirmation. This document is a draft and is not an approved meeting record.",
        size=9.5,
    )

    add_heading(doc, "Review Metadata", 1)
    add_table(
        doc,
        ["Field", "Value", "Confidence"],
        [
            ["Source", "Attached image: codex-clipboard-8ad3faf0-b95f-4232-aa0f-0bec42a7a318.png", "High"],
            ["Team", "CC", "High"],
            ["Meeting room", "C102", "High"],
            ["Recorder", "SRL", "High"],
            ["Review period", "06, 13, 20, and 28 January 2569", "High"],
            ["Document structure", "Four weekly columns with notice, follow-up, and other matters sections", "High"],
            ["Overall extraction confidence", "High for headings and dates; medium for workstreams; low for small text", "Medium"],
        ],
        [1.8, 7.1, 1.3],
    )

    add_heading(doc, "Instruction Boundary", 1)
    add_bullet(doc, "The source image is treated as information to summarize, not as an instruction to perform work.")
    add_bullet(doc, "No ClickUp update, dashboard update, task assignment, message, approval, or training registration is authorized by this document.")
    add_bullet(doc, "Unclear abbreviations, owners, deadlines, budgets, and approvals remain unverified.")

    add_heading(doc, "Notices and Reported Updates", 1)
    add_table(
        doc,
        ["Item", "Type", "Evidence in source", "Confidence"],
        [
            ["Week 3 meeting was not held because the CC team attended training", "Notice", "Week 3 notice section", "High"],
            ["Training topics included AI and Generative AI", "Training", "Week 3 notice section", "Medium"],
            ["The team was asked to keep work information current in ClickUp", "Team operating practice", "Week 2 and Week 4 notice sections", "Medium"],
            ["A dashboard was discussed as a way to view the team's work overview", "Reporting practice", "Week 4 notice section", "Medium"],
            ["A Unit Content plan connected to 15 February 2569 appears in the record", "Upcoming plan", "Week 4 notice section", "Low"],
        ],
        [2.6, 1.65, 4.75, 1.2],
    )

    add_heading(doc, "Workstream Updates", 1)
    add_table(
        doc,
        ["Workstream", "Observed content", "Status", "Source location", "Confidence"],
        [
            ["Design", "Recurring updates for Unit Design, branding, NSTIS, production, and event-related work", "Not stated", "Weeks 1, 2, and 4 follow-up sections", "Medium"],
            ["PR", "Recurring updates for public relations, MOU-related work, Run Stage, and events", "Not stated", "Weeks 1, 2, and 4 follow-up sections", "Medium"],
            ["Content", "Recurring content work including InnoStore Daily, Tech-related content, PM 2.5, and other media items", "Not stated", "Weeks 1 and 4 follow-up sections", "Low"],
            ["Studio and production", "Photo studio progress and proposed usage planning", "Needs confirmation", "Week 1 follow-up section", "Medium"],
            ["KPI and KR", "Discussion of KR and clearer individual KPI tracking", "Needs confirmation", "Week 2 notice section", "Medium"],
            ["Team operations", "Work tracking through ClickUp and dashboard reporting", "Ongoing practice", "Weeks 2 and 4 notice sections", "Medium"],
        ],
        [1.45, 5.7, 1.35, 2.15, 0.9],
    )

    add_heading(doc, "Action Register", 1)
    add_table(
        doc,
        ["Task", "Owner", "Status", "Due date", "Dependency", "Expected output", "Confidence"],
        [
            ["Keep work information current in ClickUp", "Not specified", "Reported as a team expectation", "Not specified", "Individual work updates", "Current work records", "Medium"],
            ["Clarify individual KPI and KR definitions", "Not specified", "Needs confirmation", "Not specified", "Agreed KPI and KR definitions", "KPI and KR list", "Medium"],
            ["Prepare data for dashboard reporting", "Not specified", "Needs confirmation", "Not specified", "Data structure and reporting owner", "Current team overview", "Low"],
            ["Continue weekly updates for Design, PR, and Content", "Not specified", "Recurring review item", "Weekly", "Project-specific inputs", "Updated work review", "Medium"],
            ["Attend or complete AI-related training", "CC team", "Training mentioned", "20 January 2569 appears in the source", "Training schedule", "Training participation", "Medium"],
        ],
        [2.15, 1.15, 1.55, 1.3, 2.0, 1.65, 0.8],
    )

    add_heading(doc, "Handoff Register", 1)
    add_table(
        doc,
        ["From", "To", "Required information", "Missing information", "Next action", "Confidence"],
        [
            ["CC", "PM or project team", "TOR, brief, and current work information", "Named recipient and official channel", "Confirm the handoff route", "Low"],
            ["CC", "RSP or site-related team", "Studio, space, and usage information", "Service scope and approver", "Confirm the source of current space information", "Low"],
            ["Design", "Production or event team", "Design files, media specifications, version, and production details", "Acceptance criteria and latest file", "Confirm the production handoff checklist", "Medium"],
            ["Team members", "Dashboard owner", "Work status and KPI or KR data", "Data format and update cycle", "Confirm dashboard data ownership", "Low"],
        ],
        [1.4, 1.45, 3.2, 2.8, 2.15, 0.8],
    )

    add_heading(doc, "Upcoming Dates and Events", 1)
    add_table(
        doc,
        ["Date or period", "Observed reference", "What must be confirmed", "Confidence"],
        [
            ["20 to 22 January 2569", "Date range appears in Week 2 notice content", "Which activity the date range belongs to", "Low"],
            ["20 January 2569", "Training-related date appears in Week 3 context", "Whether this is the training date or meeting date", "Medium"],
            ["15 February 2569", "Unit Content timing appears in Week 4 notice content", "Exact deliverable, owner, and approval", "Low"],
            ["May to June 2569", "Event and exhibition references appear in Week 4 follow-up content", "Event names, dates, owners, and deliverables", "Low"],
        ],
        [2.0, 4.55, 4.15, 1.0],
    )

    add_heading(doc, "Human Confirmation Required", 1)
    for item in [
        "Confirm the owners behind abbreviations such as CL, WK, CNB, WG, KHT, and other labels.",
        "Confirm which entries are decisions, which are status updates, and which are proposals.",
        "Confirm the exact project names, deadlines, and deliverables shown in small text.",
        "Confirm whether any budget figure in the image is approved, estimated, or only referenced.",
        "Confirm the official handoff route for TOR, design files, studio information, and dashboard data.",
        "Confirm whether the next version should use the original spreadsheet rather than a screenshot.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Draft Quality Result", 1)
    add_table(
        doc,
        ["Check", "Result", "Note"],
        [
            ["Section and date extraction", "Pass", "Weekly structure and main dates were identified."],
            ["Workstream grouping", "Pass with limits", "Design, PR, Content, Studio, KPI/KR, and operations were identified."],
            ["Owner and deadline safety", "Pass", "Unspecified values were not invented."],
            ["Image text extraction", "Needs improvement", "Small text requires the original spreadsheet or a higher-resolution export."],
            ["External action safety", "Pass", "No system update or task assignment was performed."],
        ],
        [2.3, 1.65, 7.0],
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(footer, "CC Team Weekly Work Review Draft", color="6B7280", size=8)

    doc.core_properties.title = "CC Team Weekly Work Review"
    doc.core_properties.subject = "Draft structured review from a CC team weekly meeting record"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "Draft document. Verify source details before operational use."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
